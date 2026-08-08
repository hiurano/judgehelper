"""
DOCX generation service for Judge Helper.
Renders markdown protocol text into native MS Word (.docx) files according to court standards.
"""
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
SIG_PATTERN = re.compile(
    r"^(Председательствующий|Секретарь(?: судебного заседания)?|Помощник судьи)(?:\s+|(?=[А-ЯЁ]))(.+)$"
)
CITY_CHECK_PATTERN = re.compile(r"^г\.\s*Нижневартовск(?:\s+.*)?$", re.IGNORECASE)
CITY_DATE_MATCH_PATTERN = re.compile(r"^(г\.\s*Нижневартовск)(?:\s+|\t+)(.+)$", re.IGNORECASE)
SPACES_TO_TAB_PATTERN = re.compile(r" {4,}")


def _parse_formatted_runs(text: str) -> list[tuple[str, bool]]:
    """Split line into chunks of (text, is_bold) based on markdown **bold**."""
    runs = []
    last_idx = 0
    for m in BOLD_PATTERN.finditer(text):
        start, end = m.span()
        if start > last_idx:
            runs.append((text[last_idx:start], False))
        runs.append((m.group(1), True))
        last_idx = end
    if last_idx < len(text):
        runs.append((text[last_idx:], False))
    return runs or [(text, False)]


def _apply_run_font(run, bold: bool = False):
    """Set font to Times New Roman 12pt with proper XML font mappings."""
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    if bold:
        run.bold = True
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(slot), "Times New Roman")


def render_docx(text: str) -> bytes:
    """
    Render protocol text as a .docx with exact Russian court document formatting:
    - Font: Times New Roman 12pt
    - Line spacing: 1.0 (Single), 0pt space before/after
    - First line indent (красная строка): 1.25 cm
    - Margins: Left 3.0 cm, Right 1.5 cm, Top 2.0 cm, Bottom 2.0 cm
    - Alignment: Centered for titles, Justified for body text
    - Right-aligned tab stop at 16.5 cm for dates, city, signatures
    - Converts markdown **bold** into native Word bold runs
    - Automatically formats signatures nicely with right-aligned tabs
    """
    doc = Document()

    # Base style: Times New Roman 12pt
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(slot), "Times New Roman")

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)

    lines = text.split("\n")
    for line in lines:
        stripped = line.strip()

        # Skip markdown code fences if LLM wrapped output
        if stripped.startswith("```"):
            continue

        # Fix collapsed signatures like "ПредседательствующийВ.А. Пономарёв" or "Помощник судьиА.И. Харитонова"
        sig_match = SIG_PATTERN.match(stripped)
        if sig_match:
            role, name = sig_match.group(1), sig_match.group(2).strip()
            line = f"{role}\t{name}"
            stripped = line.strip()

        # Handle "г. Нижневартовск [дата]" line in header — force right tab stop for date
        if CITY_CHECK_PATTERN.match(stripped):
            city_date_match = CITY_DATE_MATCH_PATTERN.match(stripped)
            if city_date_match:
                line = f"{city_date_match.group(1)}\t{city_date_match.group(2)}"
                stripped = line.strip()

        # Convert runs of 4+ spaces into tabs for clean tabular layout
        compact = SPACES_TO_TAB_PATTERN.sub("\t", line)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        # Title headers: e.g. "ПРОТОКОЛ", "судебного заседания..."
        clean_header = stripped.replace("*", "").strip()
        is_title = (
            ("ПРОТОКОЛ" in clean_header and len(clean_header) < 40)
            or clean_header.startswith("судебного заседания")
            or clean_header.startswith("по уголовному делу")
        )
        is_city_line = bool(CITY_CHECK_PATTERN.match(stripped))

        if is_title:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        else:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if is_city_line:
                p.paragraph_format.first_line_indent = Cm(0)
            elif stripped and "\t" not in compact:
                p.paragraph_format.first_line_indent = Cm(1.25)

        chunks = _parse_formatted_runs(compact)
        for chunk_text, chunk_bold in chunks:
            if not chunk_text:
                continue
            run = p.add_run(chunk_text)
            _apply_run_font(run, bold=(is_title or chunk_bold))

        if "\t" in compact:
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
