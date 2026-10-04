"""Synthetic testimony only: preservation and explicit review are the contract."""
import copy
from io import BytesIO

import pytest
from docx import Document

from backend.services.docx_generator import render_review_docx
from backend.services.protocol import (
    ReviewUpdate, apply_review, batches_for_source, corrected_utterances,
    document_blocks, document_text, normalize_source, proposals, role_registry,
    validate_proposals, sensitive_change,
)
from backend.services.review_pipeline import initial_state


@pytest.fixture
def source():
    return normalize_source({"id": "synthetic", "utterances": [
        {"speaker": "A", "text": "🙂 Я не видел видел 12 машин.", "start": 100, "end": 5000},
        {"speaker": "B", "text": "Я защитник. Суд объявляет перерыв."},
    ]})


def batch(source, **values):
    segments = batches_for_source(source)[0]
    return {"reviewed_segments": [s["id"] for s in segments],
            "corrections": [], "roles": [], "concerns": [], **values}


def correction(source, **values):
    text = source["utterances"][0]["text"]
    start = text.rindex("видел")
    return {"utterance_id": "u000001", "start": start, "end": start + 5,
            "original": "видел", "replacement": "заметил", "category": "recognition",
            "reason": "Предложение для проверки", **values}


def state_with(source, **values):
    state = initial_state()
    state["stage"] = "ready"
    result = validate_proposals(batch(source, **values), source, batches_for_source(source)[0])
    state["batches"]["0"] = {"result": result}
    return state


def test_unchanged_source_offsets_order_times_and_unicode(source):
    original = copy.deepcopy(source)
    state = state_with(source, corrections=[correction(source)])
    assert [u["text"] for u in corrected_utterances(source, state)] == [u["text"] for u in source["utterances"]]
    proposal_id = proposals(state, "corrections")[0]["id"]
    state = apply_review(source, state, ReviewUpdate(revision=0, decisions={proposal_id: "accepted"}))
    result = corrected_utterances(source, state)
    assert result[0]["text"] == "🙂 Я не видел заметил 12 машин."
    assert result[0]["start_ms"] == 100
    assert [u["id"] for u in result] == ["u000001", "u000002"]
    assert corrected_utterances(source, state) == result  # no double application
    assert source == original


@pytest.mark.parametrize("patch", [
    {"utterance_id": "missing"}, {"start": -1}, {"start": True}, {"end": 1000},
    {"original": "слышал"}, {"replacement": ""}, {"replacement": "  "},
    {"replacement": "текст\nПРОТОКОЛ"}, {"replacement": "видел"}, {"unknown": 1},
])
def test_rejects_invalid_corrections(source, patch):
    with pytest.raises(ValueError):
        validate_proposals(batch(source, corrections=[correction(source, **patch)]), source, batches_for_source(source)[0])


def test_rejects_duplicates_overlap_and_partial_processing(source):
    c = correction(source)
    for payload in (batch(source, corrections=[c, c]), batch(source, reviewed_segments=[])):
        with pytest.raises(ValueError):
            validate_proposals(payload, source, batches_for_source(source)[0])


def test_nonoverlapping_edits_apply_against_original_positions(source):
    a = correction(source, start=2, end=3, original="Я", replacement="Я лично")
    state = state_with(source, corrections=[a, correction(source)])
    state["decisions"] = {c["id"]: "accepted" for c in proposals(state, "corrections")}
    assert corrected_utterances(source, state)[0]["text"] == "🙂 Я лично не видел заметил 12 машин."


def test_large_utterance_is_covered_once_and_edits_cannot_cross_segments():
    source = normalize_source({"text": "Очень длинная реплика. " * 700})
    batches = batches_for_source(source)
    segments = [s for b in batches for s in b]
    assert "".join(s["text"] for s in segments) == source["utterances"][0]["text"]
    assert len(batches) > 1
    for a, b in zip(segments, segments[1:]):
        assert a["end"] == b["start"]
    c = {"utterance_id": "u000001", "start": segments[0]["end"] - 2,
         "end": segments[0]["end"] + 2, "original": source["utterances"][0]["text"][segments[0]["end"] - 2:segments[0]["end"] + 2],
         "replacement": "правка", "category": "spelling", "reason": "test"}
    with pytest.raises(ValueError):
        validate_proposals({"reviewed_segments": [s["id"] for s in batches[0]],
                            "corrections": [c], "roles": [], "concerns": []}, source, batches[0])


def test_roles_need_actual_evidence_and_conflicts_are_not_resolved_automatically(source):
    role = {"speaker_id": "B", "role": "Защитник", "name": "",
            "evidence": [{"utterance_id": "u000002", "quote": "Я защитник."}], "reason": "Представился"}
    state = state_with(source, roles=[role, {**role, "role": "Свидетель"}])
    assert role_registry(source, state)[1]["status"] == "conflict"
    assert corrected_utterances(source, state)[1]["label"] == "Спикер B"
    state = apply_review(source, state, ReviewUpdate(revision=0,
        speakers={"B": {"role": "Защитник", "name": ""}},
        utterances={"u000002": {"role": "Свидетель", "name": ""}}))
    assert corrected_utterances(source, state)[1]["label"] == "Свидетель"
    role["evidence"][0]["quote"] = "Такого в записи нет"
    with pytest.raises(ValueError):
        state_with(source, roles=[role])


def test_manual_text_is_traceable_and_empty_replies_cannot_disappear(source):
    state = state_with(source)
    new = apply_review(source, state, ReviewUpdate(revision=0, manual_text={"u000001": "Ручная проверка."}))
    assert corrected_utterances(source, new)[0]["original"] == source["utterances"][0]["text"]
    assert corrected_utterances(source, new)[0]["text"] == "Ручная проверка."
    for text in ("", " " * 3):
        with pytest.raises(ValueError):
            apply_review(source, state, ReviewUpdate(revision=0, manual_text={"u000001": text}))


def test_cannot_mark_unfinished_or_unreviewed_work_complete(source):
    state = state_with(source, corrections=[correction(source)], concerns=[{
        "utterance_id": "u000001", "quote": "12", "reason": "Сверить число"}])
    with pytest.raises(ValueError):
        apply_review(source, state, ReviewUpdate(revision=0, reviewed=True))
    update = ReviewUpdate(revision=0, reviewed=True,
        decisions={c["id"]: "rejected" for c in proposals(state, "corrections")},
        speakers={s: {"role": "Роль не установлена", "name": ""} for s in ("A", "B")},
        resolved_concerns=[c["id"] for c in proposals(state, "concerns")])
    assert apply_review(source, state, update)["reviewed"] is True
    state["stage"] = "error"
    with pytest.raises(ValueError):
        apply_review(source, state, update)


def test_word_uses_exact_blocks_and_only_confirmed_fields(source):
    state = state_with(source)
    blocks = document_blocks(source, state)
    text = document_text(blocks)
    assert text.count("ПРОТОКОЛ") == 1
    assert "перерыв" in text
    assert "CD-R" not in text and "Секретарь" not in text
    doc = Document(BytesIO(render_review_docx(blocks)))
    assert [p.text for p in doc.paragraphs] == [f"{b['label']}: {b['text']}" if b["kind"] == "utterance" else b["text"] for b in blocks]
    state["fields"] = {"judge": "И.И. Пример", "city": "Тестовый город"}
    blocks = document_blocks(source, state)
    assert len([b for b in blocks if b["kind"] == "signature"]) == 1
    assert "И.И. Пример" in document_text(blocks)


def test_literal_speech_is_not_parsed_as_markdown_or_a_signature():
    source = normalize_source({"text": "``` **СекретарьПример**    ПРОТОКОЛ"})
    blocks = document_blocks(source, state_with(source))
    doc = Document(BytesIO(render_review_docx(blocks)))
    assert doc.paragraphs[-1].text == "Спикер ?: " + source["utterances"][0]["text"]


@pytest.mark.parametrize('before,after', [('12', '21'), ('не видел', 'видел'), ('Иванов', 'Петров')])
def test_sensitive_edits_are_flagged_for_attention(before, after):
    assert sensitive_change({'original': before, 'replacement': after, 'category': 'recognition'})
