from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

import backend.review_api as review_api
from backend.auth import SESSION_COOKIE, make_session_token
from backend.db import jobs, user_store
from backend.main import app
from backend.services.protocol import normalize_source
from backend.services.review_pipeline import initial_state


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(review_api, 'spawn', lambda coro, name: coro.close())
    user_store.create_user('review-owner', 'Test-password-123')
    user_store.create_user('review-other', 'Test-password-456')
    with TestClient(app) as client:
        source = normalize_source({'text': 'Исходная тестовая реплика.'})
        state = initial_state()
        state['stage'] = 'ready'
        jobs['review-api'] = {'user_id': 'review-owner', 'status': 'done', 'protocol_mode': 'review', 'has_review': True}
        jobs.create_review('review-api', source, state)
        client.cookies.set(SESSION_COOKIE, make_session_token('review-owner'))
        yield client
        jobs.delete_for_user('review-owner')
        jobs.delete_for_user('review-other')


def test_owner_access_and_stale_revision_do_not_overwrite_original(client):
    record = client.get('/jobs/review-api/review').json()
    assert record['source']['utterances'][0]['text'] == 'Исходная тестовая реплика.'
    payload = {'revision': record['revision'], 'manual_text': {'u000001': 'Уточнённая реплика.'}}
    saved = client.put('/jobs/review-api/review', json=payload)
    assert saved.status_code == 200
    assert saved.json()['rows'][0]['text'] == 'Уточнённая реплика.'
    assert saved.json()['source'] == record['source']
    assert client.put('/jobs/review-api/review', json=payload).status_code == 409
    assert jobs.get_review('review-api')['source'] == record['source']
    assert 'source' not in client.get('/status/review-api').json()
    assert 'source' not in client.get('/jobs').json()['jobs'][0]


@pytest.mark.parametrize('method,path', [
    ('get', '/jobs/review-api/review'), ('get', '/jobs/review-api/review/docx'),
    ('put', '/jobs/review-api/review'), ('post', '/jobs/review-api/review/retry'),
])
def test_every_review_route_is_owner_only(client, method, path):
    call = getattr(client, method)
    kwargs = {'json': {'revision': 0}} if method == 'put' else {}
    client.cookies.set(SESSION_COOKIE, make_session_token('review-other'))
    assert call(path, **kwargs).status_code == 404
    client.cookies.clear()
    assert call(path, **kwargs).status_code == 401


def test_export_comes_from_saved_blocks_and_preserves_literal_text(client):
    payload = {'revision': 0, 'manual_text': {'u000001': '**Буквальный** текст.'}}
    assert client.put('/jobs/review-api/review', json=payload).status_code == 200
    response = client.get('/jobs/review-api/review/docx')
    assert response.status_code == 200
    document = Document(BytesIO(response.content))
    assert document.paragraphs[-1].text == 'Спикер ?: **Буквальный** текст.'
    assert document.paragraphs[0].text.startswith('ЧЕРНОВИК')
    record = client.get('/jobs/review-api/review').json()
    payload.update(revision=record['revision'], reviewed=True,
                   speakers={'?': {'role': 'Роль не установлена', 'name': ''}})
    assert client.put('/jobs/review-api/review', json=payload).status_code == 200
    assert jobs.get('review-api')['review_status'] == 'reviewed'
    document = Document(BytesIO(client.get('/jobs/review-api/review/docx').content))
    assert document.paragraphs[0].text == 'ПРОТОКОЛ'


def test_cannot_replace_source_add_fake_replies_or_review_during_analysis(client):
    assert client.put('/jobs/review-api/review', json={'revision': 0, 'source': {}}).status_code == 422
    assert client.put('/jobs/review-api/review', json={'revision': 0, 'manual_text': {'fake': 'text'}}).status_code == 400
    record = jobs.get_review('review-api')
    record['state']['stage'] = 'processing'
    jobs.save_review('review-api', record['state'], 0)
    assert client.put('/jobs/review-api/review', json={'revision': 1}).status_code == 400


def test_retry_has_atomic_slot_limit_and_does_not_discard_source(client, monkeypatch):
    record = jobs.get_review('review-api')
    record['state']['stage'] = 'error'
    jobs.save_review('review-api', record['state'], 0, {'status': 'error'})
    monkeypatch.setattr(review_api, 'MAX_ACTIVE_JOBS_PER_USER', 1)
    jobs['review-busy'] = {'user_id': 'review-owner', 'status': 'processing'}
    assert client.post('/jobs/review-api/review/retry').status_code == 429
    assert jobs.get_review('review-api')['state']['stage'] == 'error'
    jobs.delete('review-busy')
    assert client.post('/jobs/review-api/review/retry').status_code == 200
    assert jobs.get('review-api')['status'] == 'processing'
    assert jobs.get_review('review-api')['source'] == record['source']
    assert client.post('/jobs/review-api/review/retry').status_code == 409


def test_delete_removes_artifacts_and_legacy_rows_are_unchanged(client):
    jobs['review-legacy'] = {'user_id': 'review-owner', 'status': 'done', 'draft': 'Старый протокол'}
    assert client.get('/jobs/review-legacy/review').status_code == 404
    assert client.get('/status/review-legacy').json()['draft'] == 'Старый протокол'
    assert client.delete('/jobs/review-api').status_code == 200
    assert jobs.get_review('review-api') is None
    assert client.get('/jobs/review-api/review').status_code == 404


def test_new_upload_captures_mode_and_late_webhook_cannot_change_review(client, monkeypatch):
    import backend.main as main
    monkeypatch.setattr(main, 'PROTOCOL_MODE', 'review')
    job_id = client.post('/uploads').json()['job_id']
    assert jobs.get(job_id)['protocol_mode'] == 'review'
    monkeypatch.setattr(main, 'PROTOCOL_MODE', 'legacy')
    assert jobs.get(job_id)['protocol_mode'] == 'review'
    jobs.delete(job_id)
    monkeypatch.setattr(main, 'WEBHOOK_SECRET', 'test-webhook-secret')
    response = client.post('/webhook/aai', json={'transcript_id': 'review-api', 'status': 'error'},
                           headers={'x-webhook-secret': 'test-webhook-secret'})
    assert response.status_code == 200
    assert jobs.get('review-api')['status'] == 'done'
    assert jobs.get_review('review-api')['state']['stage'] == 'ready'
